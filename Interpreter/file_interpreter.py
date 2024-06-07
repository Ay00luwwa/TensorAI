import PyPDF2
import openpyxl
import pandas as pd
from docx import Document

def read_pdf(file_path):
    with open(file_path, 'rb') as file:
        reader = PyPDF2.PdfFileReader(file)
        content = ''
        for page in range(reader.numPages):
            content += reader.getPage(page).extract_text()
    return content

def read_xlsx(file_path):
    workbook = openpyxl.load_workbook(file_path)
    sheet = workbook.active
    content = ''
    for row in sheet.iter_rows(values_only=True):
        content += '\t'.join([str(cell) for cell in row]) + '\n'
    return content

def read_csv(file_path):
    df = pd.read_csv(file_path)
    return df.to_string()

def read_docx(file_path):
    doc = Document(file_path)
    content = '\n'.join([para.text for para in doc.paragraphs])
    return content
